import sys
import os
import json
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

def get_color_hex(color_obj):
    try:
        if color_obj and color_obj.type == 1: # RGB Color
            rgb = color_obj.rgb
            return f"#{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}"
    except Exception:
        pass
    return None

def parse_pptx(pptx_path):
    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        print(f"ERROR: Could not load PPTX presentation: {e}")
        sys.exit(1)
        
    base_name = os.path.basename(pptx_path)
    output_dir = os.path.join(os.path.dirname(pptx_path), "extracted", base_name)
    os.makedirs(output_dir, exist_ok=True)
    
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    
    presentation_data = {
        "title": base_name,
        "width": slide_width,
        "height": slide_height,
        "slides": []
    }
    
    for i, slide in enumerate(prs.slides):
        slide_data = {
            "slideNumber": i + 1,
            "background": None,
            "elements": []
        }
        
        # Try to parse slide background solid color
        try:
            background = slide.background
            if background and background.fill.type == 1: # Solid
                bg_color = get_color_hex(background.fill.fore_color)
                if bg_color:
                    slide_data["background"] = bg_color
        except Exception:
            pass
            
        for shape in slide.shapes:
            # Common positioning percentages
            try:
                left_pct = (shape.left / slide_width) * 100
                top_pct = (shape.top / slide_height) * 100
                width_pct = (shape.width / slide_width) * 100
                height_pct = (shape.height / slide_height) * 100
            except Exception:
                continue
                
            element = {
                "name": shape.name,
                "left": left_pct,
                "top": top_pct,
                "width": width_pct,
                "height": height_pct,
                "fill_color": None,
                "border_color": None,
                "type": "shape"
            }
            
            # Solid shape fill color
            try:
                if shape.fill.type == 1: # Solid fill
                    fill_hex = get_color_hex(shape.fill.fore_color)
                    if fill_hex:
                        element["fill_color"] = fill_hex
            except Exception:
                pass
                
            # Line border color
            try:
                if shape.line and shape.line.color:
                    line_hex = get_color_hex(shape.line.color)
                    if line_hex:
                        element["border_color"] = line_hex
            except Exception:
                pass
                
            # Image element extraction
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                    image_bytes = image.blob
                    image_ext = image.ext
                    img_name = f"slide_{i+1}_{shape.shape_id}.{image_ext}"
                    img_path = os.path.join(output_dir, img_name)
                    
                    with open(img_path, "wb") as f:
                        f.write(image_bytes)
                        
                    element["type"] = "image"
                    element["src"] = f"../uploads/resources/extracted/{base_name}/{img_name}"
                except Exception as img_err:
                    print(f"Warning parsing picture: {img_err}")
                    
            # Text Frame parsing
            if shape.has_text_frame:
                element["type"] = "text"
                element["paragraphs"] = []
                
                for paragraph in shape.text_frame.paragraphs:
                    p_data = {
                        "text": paragraph.text,
                        "align": "left",
                        "runs": []
                    }
                    
                    # Determine paragraph alignment
                    if paragraph.alignment == PP_ALIGN.CENTER:
                        p_data["align"] = "center"
                    elif paragraph.alignment == PP_ALIGN.RIGHT:
                        p_data["align"] = "right"
                    elif paragraph.alignment == PP_ALIGN.JUSTIFY:
                        p_data["align"] = "justify"
                        
                    for run in paragraph.runs:
                        run_data = {
                            "text": run.text,
                            "bold": bool(run.font.bold),
                            "italic": bool(run.font.italic),
                            "underline": bool(run.font.underline),
                            "color": get_color_hex(run.font.color),
                            "size": run.font.size.pt if run.font.size else None,
                            "font_name": run.font.name
                        }
                        p_data["runs"].append(run_data)
                        
                    element["paragraphs"].append(p_data)
                    
                # Store full text representation for remote or simple fallbacks
                element["text"] = shape.text.strip()
                
            slide_data["elements"].append(element)
            
        presentation_data["slides"].append(slide_data)
        
    json_path = pptx_path + ".json"
    with open(json_path, "w") as f:
        json.dump(presentation_data, f, indent=4)
    print(f"SUCCESS: Parsed {len(prs.slides)} slides with detailed coordinates.")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python parse_pptx.py <path_to_pptx>")
    else:
        parse_pptx(sys.argv[1])
